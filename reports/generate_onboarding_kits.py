"""Automated Onboarding Kit Generator (DOCX & PDF) for all 8 Tenant Verticals.

Generates pre-defined, professionally formatted onboarding kits in both:
  1. Editable DOCX (.docx) with structured tables and clear section headers
  2. Standalone clean PDF (.pdf) matching the specifications

The generated documents serve as templates given to clients to fill out and upload.
The backend's Section-by-Section Targeted Ollama Bifurcator identifies the standardized
section headers to parse brand, menu, catalog/offerings, FAQs, and bot guardrails.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Dict, List

import docx
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
import fitz  # PyMuPDF

ROOT_DIR = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT_DIR / "backend" / "static" / "templates" / "onboarding"

# Ensure output directory exists
OUT_DIR.mkdir(parents=True, exist_ok=True)

# 8 Supported Verticals
VERTICAL_CONFIGS: Dict[str, Dict[str, Any]] = {
    "real_estate": {
        "title": "Real Estate & Property Development",
        "slug": "real_estate",
        "tagline": "Residences, commercial spaces, floor plans, pricing & site visits",
        "default_company": "Skyline Realty Group",
        "default_tenant": "skylinerealty",
        "default_bot": "Skyline Advisor",
        "default_hours": "09:30 AM - 07:00 PM (Mon-Sat)",
        "menu_items": [
            ("Explore Properties", "2, 3 & 4 BHK luxury residences & penthouses", "🏢 Properties", "offerings"),
            ("Ongoing Projects", "Prime locations, masterplans & tower specs", "🏢 Properties", "offering_details"),
            ("Schedule Site Visit", "Free cab pickup & guided property tour", "📞 Bookings", "book_callback"),
            ("Price & Payment Plans", "Cost sheets, EMI calculations & construction plans", "📞 Bookings", "quote"),
            ("RERA & Legal Sanctions", "RERA registration, sanctions & bank approvals", "ℹ️ Legal", "faq"),
            ("Customer Support", "Existing buyer assistance & documentation help", "🆘 Support", "raise_ticket"),
        ],
        "catalog_columns": [
            "Property Name", "Type", "Configuration", "Carpet Area (sqft)",
            "Price Range", "Location", "Key Amenities", "Possession", "RERA ID"
        ],
        "catalog_rows": [
            [
                "Skyline Elegance", "Apartment", "2 & 3 BHK", "1,250 - 1,680",
                "₹85 Lakhs - ₹1.25 Cr", "Prime Avenue, Whitefield",
                "Clubhouse, Swimming Pool, Gym, Tennis Court", "Ready to Move", "PRM/KA/RERA/1251/2023"
            ],
            [
                "Skyline Grandeur", "Luxury Penthouse", "4 BHK Duplex", "3,400 - 4,200",
                "₹2.8 Cr - ₹3.6 Cr", "Indiranagar 100ft Rd",
                "Private Deck, Sky Lounge, Concierge, 3 Car Parks", "Dec 2026", "PRM/KA/RERA/1251/2024"
            ],
            [
                "Skyline Horizon Tech Park", "Commercial Office", "Grade-A Bare Shell", "2,500 - 25,000",
                "₹95 / sq.ft. (Lease)", "Outer Ring Road (ORR)",
                "100% Power Backup, LEED Gold, High-Speed Elevators", "Immediate", "PRM/KA/RERA/1251/2022"
            ],
            [
                "Skyline Orchard Green", "Gated Villa", "4 BHK Luxury Villa", "2,850 sq.ft.",
                "₹1.95 Cr", "Sarjapur Road",
                "Private Garden, Solar Powered, Clubhouse & Lake View", "June 2027", "PRM/KA/RERA/1251/2025"
            ],
        ],
        "faq_rows": [
            ("What banks have approved home loans for these projects?", "Approved by SBI, HDFC, ICICI, and Axis Bank with instant sanction and up to 80% financing.", "Home Loans"),
            ("Is the site visit pickup complimentary?", "Yes, we arrange a dedicated AC cab pickup and drop from your residence or office at zero charge.", "Site Visits"),
            ("What is the booking advance amount?", "A booking advance of ₹2,00,000 secures your selected unit. The remainder is milestone-linked.", "Booking & Payments"),
            ("Are all projects RERA certified?", "Yes, every project holds approved sanctions from the local municipal authority and valid RERA certification.", "Legal & Approvals"),
            ("Can Non-Resident Indians (NRIs) purchase units?", "Yes, NRIs can buy residential or commercial units using standard NRE/NRO banking channels with full RBI compliance.", "NRI Purchase"),
        ],
    },
    "ecommerce": {
        "title": "E-Commerce & Retail Brands",
        "slug": "ecommerce",
        "tagline": "Catalog, shopping cart, orders, instant checkout & return tracking",
        "default_company": "Urban Luxe Lifestyle",
        "default_tenant": "urbanluxe",
        "default_bot": "Luxe Concierge",
        "default_hours": "24/7 Automated Support (Office: 09:00 AM - 06:00 PM)",
        "menu_items": [
            ("Browse Catalog", "Explore latest collections, deals & new arrivals", "🛍️ Shop", "offerings"),
            ("Track My Order", "Real-time delivery status & courier tracking", "📦 Orders", "track_order"),
            ("Returns & Refunds", "Initiate return or replacement within 7 days", "📦 Orders", "returns"),
            ("Exclusive Offers", "Discount coupons & seasonal festival promos", "🏷️ Deals", "faq"),
            ("Customer Support", "Chat with a support agent regarding order queries", "🆘 Help", "raise_ticket"),
        ],
        "catalog_columns": [
            "Product Name", "SKU", "Category", "Price (₹)", "Discount Price",
            "Key Specifications / Size", "Stock Status", "Warranty / Returns"
        ],
        "catalog_rows": [
            [
                "Italian Leather Briefcase", "UL-BAG-001", "Accessories", "₹4,999", "₹3,799",
                "Full-grain calf leather, 15.6 inch laptop sleeve, Brown", "In Stock (45)", "1 Year Warranty / 7-Day Return"
            ],
            [
                "Ultra-Quiet Sonic Toothbrush", "UL-CARE-010", "Personal Care", "₹2,499", "₹1,899",
                "40,000 VPM, 4 Brush Heads, 30-day battery life, Waterproof", "In Stock (120)", "2 Year Replacement Warranty"
            ],
            [
                "Organic Bamboo Cotton Sheets", "UL-HOME-045", "Home & Living", "₹3,499", "₹2,999",
                "King size, 400 TC, Breathable hypoallergenic fabric, Ivory", "In Stock (80)", "7-Day Easy Return"
            ],
            [
                "Minimalist Chrono Watch", "UL-WAT-102", "Watches", "₹7,999", "₹5,999",
                "Sapphire crystal, Japanese quartz movement, 5ATM water resist", "Low Stock (8)", "2 Year International Warranty"
            ],
        ],
        "faq_rows": [
            ("What is your standard delivery timeline?", "Metro cities receive orders in 2-3 business days. Other regions take 4-6 business days with live tracking.", "Shipping"),
            ("How do I return a product?", "Tap 'Returns & Refunds' or message your Order ID. We arrange free courier pickup within 48 hours.", "Returns"),
            ("What payment methods do you accept?", "We accept UPI (GPay, PhonePe, Paytm), Credit/Debit Cards, Net Banking, and Cash on Delivery (COD).", "Payments"),
            ("How do I track my active order?", "Type 'Track Order' followed by your 6-digit Order ID or registered mobile number to get the live courier link.", "Tracking"),
        ],
    },
    "it_software": {
        "title": "Software & IT Services",
        "slug": "it_software",
        "tagline": "Software engineering, cloud architecture, tech stack & project quotes",
        "default_company": "Apex Cloud & Software Solutions",
        "default_tenant": "apexsoftware",
        "default_bot": "Apex Tech Advisor",
        "default_hours": "09:00 AM - 06:30 PM (Mon-Fri)",
        "menu_items": [
            ("Our Services", "Custom web, mobile apps, cloud & AI engineering", "💻 Services", "offerings"),
            ("Tech Stack", "React, Node.js, Python, AWS, Flutter & Kubernetes", "💻 Services", "offering_details"),
            ("Get Project Quote", "Tell us your requirements for an instant estimate", "📋 Estimates", "quote"),
            ("Schedule Tech Call", "30-minute discovery call with our lead architect", "📋 Estimates", "book_callback"),
            ("Case Studies", "Delivered enterprise projects & client testimonials", "ℹ️ Portfolio", "faq"),
        ],
        "catalog_columns": [
            "Service / Offering", "Category", "Starting Price / Model",
            "Tech Stack & Tools", "Typical Timeline", "Scope & Deliverables"
        ],
        "catalog_rows": [
            [
                "Full-Stack Web App Development", "Custom Software", "Starting $5,000 / Fixed or $45/hr",
                "Next.js, FastAPI, PostgreSQL, AWS, Docker", "6 - 12 Weeks",
                "End-to-end architecture, responsive UI, REST/GraphQL APIs, CI/CD pipeline"
            ],
            [
                "Cross-Platform Mobile App", "Mobile Engineering", "Starting $6,500 / Project",
                "Flutter / React Native, Firebase, Supabase", "8 - 14 Weeks",
                "iOS & Android build, push notifications, offline cache, App Store publishing"
            ],
            [
                "Cloud Architecture & DevOps", "Cloud Engineering", "Starting $2,500 / Month Retainer",
                "AWS, GCP, Terraform, Kubernetes, Helm", "Ongoing / SOW",
                "Infrastructure as Code, zero-downtime deployment, auto-scaling, cost optimization"
            ],
            [
                "Enterprise AI & LLM Integrations", "AI Engineering", "Starting $8,000 / Project",
                "OpenAI, Ollama, LangChain, FAISS, Python", "4 - 8 Weeks",
                "Private RAG knowledge retrieval, vector search, customer AI chatbot bots"
            ],
        ],
        "faq_rows": [
            ("How do you estimate project costs?", "We evaluate your feature list, user roles, APIs, and timeline, providing fixed-bid or dedicated agile squad models.", "Pricing"),
            ("Do you sign an NDA before discussions?", "Yes, we execute mutual non-disclosure agreements (NDAs) before reviewing any client architecture or IP.", "Security & IP"),
            ("Who owns the intellectual property (IP)?", "The client holds 100% ownership of source code, designs, and data upon milestone payment settlement.", "IP Ownership"),
            ("What post-launch maintenance do you provide?", "Every project includes 60 days of complimentary bug-fixing and warranty support, with optional SLAs.", "Support"),
        ],
    },
    "tours_travel": {
        "title": "Tours & Travel Agency",
        "slug": "tours_travel",
        "tagline": "Holiday packages, curated itineraries, hotel bookings & trip quotes",
        "default_company": "Wanderly Voyages & Expeditions",
        "default_tenant": "wanderly",
        "default_bot": "Wanderly Trip Guide",
        "default_hours": "09:00 AM - 08:00 PM (All 7 Days)",
        "menu_items": [
            ("Holiday Packages", "Domestic & international curated holiday packages", "✈️ Packages", "offerings"),
            ("Custom Itinerary", "Build a personalized honeymoon, family or group trip", "✈️ Packages", "quote"),
            ("Visa & Documents", "Tourist visa requirements, passport & travel insurance", "📋 Travel Info", "faq"),
            ("Speak to Travel Expert", "Call our destination specialist for flight & hotel advice", "📞 Contact", "book_callback"),
            ("Cancellation Policy", "Refund terms, reschedule rules & trip insurance", "ℹ️ Policies", "faq"),
        ],
        "catalog_columns": [
            "Package Name", "Destination", "Duration", "Price per Person",
            "Key Inclusions", "Highlights", "Best Travel Months"
        ],
        "catalog_rows": [
            [
                "Bali Island Escape", "Bali, Indonesia", "6 Nights / 7 Days", "₹42,999 + Flight",
                "4-Star Private Pool Villa, Breakfast, AC Transfers, Water Sports", "Ubud rice terraces, Nusa Penida tour, Tanah Lot sunset", "April to October"
            ],
            [
                "Magical Kashmir Valley", "Srinagar & Gulmarg", "5 Nights / 6 Days", "₹26,500 per person",
                "Deluxe Dal Lake Houseboat, 4-Star Hotels, Gondola Pass, All Dinners", "Shikara ride, Gulmarg snow point, Pahalgam valley tour", "Oct to March (Snow) / Apr to Jun"
            ],
            [
                "Swiss Alps & Paris Romance", "France & Switzerland", "8 Nights / 9 Days", "₹1,45,000 per person",
                "Eurail Pass, Central 4-Star Hotels, Eiffel Tower Pass, Mt. Titlis", "Paris Seine cruise, Lucerne bridge, Jungfraujoch top of Europe", "May to September"
            ],
            [
                "Dubai Extravaganza", "Dubai, UAE", "4 Nights / 5 Days", "₹32,000 + Visa",
                "4-Star Hotel, Desert Safari with BBQ, Dhow Cruise, Burj Khalifa 124th", "Dubai Mall fountain, Miracle Garden, Marina yacht cruise", "November to March"
            ],
        ],
        "faq_rows": [
            ("Are flights included in the package cost?", "Packages can be booked with or without flights based on your preferred departure city and schedule.", "Flights"),
            ("What is your cancellation and refund policy?", "Cancellations made 21 days before departure receive a 90% refund. Rescheduling is available with no penalty.", "Refunds"),
            ("Do you assist with tourist visas?", "Yes, our in-house visa concierge assists with document vetting, appointment booking, and visa processing.", "Visa"),
            ("Can packages be customized for families?", "Absolutely! Every package can be tailored for room configurations, dietary needs, or extra days.", "Customization"),
        ],
    },
    "healthcare": {
        "title": "Healthcare & Specialist Clinics",
        "slug": "healthcare",
        "tagline": "Specialties, doctor consultation timings, health packages & appointments",
        "default_company": "CareWell Multi-Specialty Clinic",
        "default_tenant": "carewellclinic",
        "default_bot": "CareWell Care Assistant",
        "default_hours": "08:00 AM - 08:30 PM (Mon-Sat), Emergency 24/7",
        "menu_items": [
            ("Our Specialties", "Cardiology, Orthopedics, Pediatrics, Dermatology & Dental", "🩺 Medical", "offerings"),
            ("Book Appointment", "Schedule in-person clinic or video consultation", "🩺 Medical", "book_callback"),
            ("Health Checkup Packages", "Comprehensive annual, senior & executive wellness packages", "💊 Wellness", "offering_details"),
            ("Lab Tests & Reports", "Blood tests, radiology timings & digital test report access", "🔬 Diagnostics", "faq"),
            ("Emergency & Ambulance", "24x7 emergency helpline & fast ambulance dispatch", "🆘 Emergency", "faq"),
        ],
        "catalog_columns": [
            "Specialty / Department", "Senior Specialist Doctor", "Consultation Fee",
            "Available Timings", "Services & Treatments", "Room / Clinic wing"
        ],
        "catalog_rows": [
            [
                "Cardiology & Heart Care", "Dr. Rajesh Sharma, MD, DM", "₹900",
                "Mon, Wed, Fri (10:00 AM - 02:00 PM)", "ECG, 2D Echo, TMT, Hypertension management", "Wing A, Room 104"
            ],
            [
                "Orthopedics & Joint Surgery", "Dr. Sunita Rao, MS (Ortho)", "₹800",
                "Tue, Thu, Sat (02:00 PM - 06:00 PM)", "Joint pain, arthritis, fracture care, sports rehab", "Wing B, Room 201"
            ],
            [
                "Pediatrics & Child Health", "Dr. Amit Verma, DCH, DNB", "₹700",
                "Mon to Sat (09:00 AM - 01:00 PM)", "Newborn vaccination, child growth monitoring, allergy care", "Wing A, Room 108"
            ],
            [
                "Dermatology & Cosmetology", "Dr. Neha Kulkarni, MD", "₹850",
                "Mon, Wed, Sat (04:00 PM - 08:00 PM)", "Acne treatment, laser therapy, hair loss PRP, skin allergy", "Wing C, Room 302"
            ],
        ],
        "faq_rows": [
            ("How do I book an appointment with a doctor?", "Tap 'Book Appointment' or message the doctor's name or department to choose an open slot.", "Appointments"),
            ("How do I receive my diagnostic lab test reports?", "Lab reports are automatically sent to your WhatsApp number as downloadable PDF files within 12 hours.", "Lab Reports"),
            ("Do you accept insurance and cashless claims?", "Yes, we provide cashless coverage for all major TPA and corporate medical insurance providers.", "Insurance"),
            ("What should I do in case of a medical emergency?", "Please call our 24/7 emergency hotline immediately at +91 9988776655 for immediate triage.", "Emergency"),
        ],
    },
    "banking": {
        "title": "Banking & Insurance Services",
        "slug": "banking",
        "tagline": "Accounts, deposits, credit cards, claim processes & branch locator",
        "default_company": "Meridian Premier Bank",
        "default_tenant": "meridianbank",
        "default_bot": "Meridian Assistant",
        "default_hours": "09:30 AM - 05:00 PM (Mon-Fri, 1st/3rd/5th Sat)",
        "menu_items": [
            ("Savings & Deposits", "Zero-balance savings, high-interest fixed deposits & recurring", "💳 Accounts", "offerings"),
            ("Credit Cards & Loans", "Lifestyle cashback cards, personal & vehicle loans", "💳 Accounts", "offering_details"),
            ("Insurance & Claims", "Life, health & motor insurance claim settlement guidelines", "🛡️ Insurance", "faq"),
            ("Locate Branch & ATM", "Find nearby branch locations, working hours & IFSC codes", "ℹ️ Info", "faq"),
            ("Block Lost Card", "Instantly block stolen or misplaced debit/credit card", "🆘 Security", "raise_ticket"),
        ],
        "catalog_columns": [
            "Product / Scheme", "Category", "Interest Rate / Annual Fee",
            "Key Features & Benefits", "Eligibility Criteria", "Minimum Balance"
        ],
        "catalog_rows": [
            [
                "Meridian Elite Savings Account", "Savings", "Up to 7.10% p.a.",
                "Free platinum debit card, unlimited ATM withdrawals, airport lounge access", "Resident Indian, 18+ years", "₹10,000 AMB"
            ],
            [
                "Smart Fixed Deposit Scheme", "Deposits", "7.65% p.a. (8.15% for Seniors)",
                "Flexible tenure (7 days to 10 yrs), monthly payout option, loan against FD", "Individual or Joint", "Min ₹5,000"
            ],
            [
                "Meridian CashBack Credit Card", "Credit Card", "₹499/yr (Waived on ₹50k spend)",
                "5% cashback on online food & shopping, 1% fuel surcharge waiver", "Annual Income > ₹3.6 Lakhs", "N/A"
            ],
            [
                "Express Home Loan", "Loans", "Starting 8.35% p.a.",
                "Zero prepayment penalty, quick sanction in 48 hours, up to 30 year tenure", "Salaried / Self-Employed", "N/A"
            ],
        ],
        "faq_rows": [
            ("How do I immediately block a lost or stolen card?", "Type 'Block Card' or tap 'Block Lost Card'. You will receive an SMS confirmation within 10 seconds.", "Card Security"),
            ("Can financial transactions be made through WhatsApp?", "For your safety, our chatbot provides answers, eligibility checks, and statements only — never money transfers.", "Security Guardrail"),
            ("What documents are required to open a savings account?", "Aadhaar Card, PAN Card, and a passport-size photo. Video KYC is completed online in 5 minutes.", "Account Opening"),
            ("How do I submit an insurance claim?", "Tap 'Insurance & Claims' for step-by-step documentation guidelines and instant hospital TPA approval lists.", "Claims"),
        ],
    },
    "finance": {
        "title": "Finance, Accounting & Tax Advisory",
        "slug": "finance",
        "tagline": "Taxation, corporate audit, compliance, payroll & virtual CFO services",
        "default_company": "Northwind Financial Advisors",
        "default_tenant": "northwindfin",
        "default_bot": "Northwind Tax Advisor",
        "default_hours": "09:30 AM - 06:30 PM (Mon-Sat)",
        "menu_items": [
            ("Advisory Services", "Income tax, corporate audit, GST compliance & virtual CFO", "📊 Services", "offerings"),
            ("Book Tax Consultation", "Schedule 45-minute consultation with a Chartered Accountant", "📅 Consultation", "book_callback"),
            ("Pricing & Retainers", "Monthly retainer plans for startups, SMEs & enterprises", "💼 Pricing", "quote"),
            ("Tax Deadlines & Guides", "Upcoming GST, advance tax, TDS & MCA filing calendar", "📅 Compliance", "faq"),
            ("Upload Client Documents", "Secure document submission portal for audits & tax filing", "📁 Portal", "raise_ticket"),
        ],
        "catalog_columns": [
            "Service Name", "Category", "Fee Structure / Pricing",
            "Deliverables & Scope", "Turnaround Time", "Target Client"
        ],
        "catalog_rows": [
            [
                "Individual & NRI Tax Filing (ITR)", "Direct Taxation", "Starting ₹1,999 per return",
                "Capital gains calculation, foreign income disclosure, tax refund filing", "2 - 4 Business Days", "Salaried, Professionals, NRIs"
            ],
            [
                "Monthly GST Compliance & Filing", "Indirect Taxation", "Starting ₹3,500 / month",
                "GSTR-1, GSTR-3B filing, input tax credit reconciliation, notice response", "Monthly by 10th/20th", "Proprietorships, LLPs, Pvt Ltd"
            ],
            [
                "Startup Virtual CFO & Accounting", "Corporate Advisory", "Starting ₹18,000 / month",
                "Bookkeeping in QuickBooks/Zoho, MIS reports, payroll, board deck preparation", "Continuous Retainer", "Startups & High-growth SMEs"
            ],
            [
                "Statutory & Internal Audit", "Audit & Assurance", "Starting ₹25,000 / audit",
                "Comprehensive balance sheet audit, internal control review, CARO reporting", "2 - 3 Weeks", "Private & Public Companies"
            ],
        ],
        "faq_rows": [
            ("What documents do I need to share for Income Tax Filing?", "Form 16, Form 26AS, AIS/TIS summary, bank interest statements, and investment proofs.", "ITR Documents"),
            ("How do you ensure data confidentiality?", "All financial statements are encrypted with AES-256 and protected by non-disclosure agreements.", "Confidentiality"),
            ("Do you handle GST notices and scrutiny assessments?", "Yes, our team drafts legal responses and represents clients before tax authorities.", "Notices & Scrutiny"),
            ("Can you manage multi-state GST registrations?", "Yes, we handle pan-India GST registrations, compliance, and input credit optimization across all states.", "Multi-State GST"),
        ],
    },
    "generic": {
        "title": "General Business & Enterprise Services",
        "slug": "generic",
        "tagline": "Company overview, offerings, pricing, interactive menu & 24/7 automated FAQ support",
        "default_company": "Acme Global Solutions",
        "default_tenant": "acmeglobal",
        "default_bot": "Acme Assistant",
        "default_hours": "09:00 AM - 06:00 PM (Mon-Fri)",
        "menu_items": [
            ("Our Offerings", "Browse all products, packages and specialized services", "📦 Offerings", "offerings"),
            ("Pricing & Quote", "Request pricing estimates tailored to your needs", "💼 Pricing", "quote"),
            ("Schedule a Callback", "Connect directly with our representative", "📞 Contact", "book_callback"),
            ("Knowledge & FAQs", "Common questions, delivery terms and company policies", "ℹ️ FAQs", "faq"),
            ("Customer Support", "Raise a ticket or speak with a live support executive", "🆘 Support", "raise_ticket"),
        ],
        "catalog_columns": [
            "Item / Service Name", "Category", "Price / Rate", "Unit",
            "Key Features & Inclusions", "Detailed Description"
        ],
        "catalog_rows": [
            [
                "Standard Business Tier", "Core Offerings", "₹9,999", "per month",
                "Includes all essential features, standard support, up to 5 user seats", "Ideal for growing teams looking to establish baseline operations"
            ],
            [
                "Professional Enterprise Tier", "Core Offerings", "₹24,999", "per month",
                "Unlimited user seats, 24/7 SLA, priority account manager, custom integrations", "Complete end-to-end suite for scaling companies"
            ],
            [
                "Custom Consultation & Setup", "Professional Services", "Starting ₹15,000", "one-time",
                "Architecture audit, configuration, team onboarding & training session", "Personalized onboarding led by our senior technical staff"
            ],
        ],
        "faq_rows": [
            ("What are your business operating hours?", "Our digital bot operates 24/7. Human team support is active Monday through Friday, 9:00 AM to 6:00 PM.", "Hours"),
            ("How can I request a customized quotation?", "Tap 'Pricing & Quote' or describe your requirements in the chat to receive a detailed cost breakdown.", "Quotes"),
            ("What is your refund or cancellation policy?", "Full refunds are provided within 14 days of subscription onboarding if you are not fully satisfied.", "Refunds"),
            ("How do I contact a senior representative?", "Tap 'Schedule a Callback' and select your preferred time slot; we will call you directly.", "Contact"),
        ],
    },
}


# ──────────────────────────────────────────────────────────────────────────────
# 1. DOCX Generation Helper Functions
# ──────────────────────────────────────────────────────────────────────────────

def _set_cell_background(cell, hex_color: str):
    """Set cell background shading in docx table."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_table_borders(table, color="D3D3D3"):
    """Set clean light borders on docx table."""
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)


def generate_docx_kit(vertical_key: str, config: Dict[str, Any], output_path: Path):
    """Generate an editable, cleanly structured .docx Onboarding Kit."""
    doc = Document()

    # Set page margins
    sections = doc.sections
    for sec in sections:
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.75)
        sec.left_margin = Inches(0.75)
        sec.right_margin = Inches(0.75)

    # Document Header
    title_p = doc.add_paragraph()
    title_run = title_p.add_run(f"📋 AI WhatsApp Chatbot Onboarding Kit\n{config['title']}")
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 41, 59)  # Slate-800
    title_p.paragraph_format.space_after = Pt(4)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run(f"Industry Vertical: {config['slug']} | {config['tagline']}")
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)  # Slate-500
    sub_p.paragraph_format.space_after = Pt(14)

    # Instructions Box
    instr_table = doc.add_table(rows=1, cols=1)
    instr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = instr_table.cell(0, 0)
    _set_cell_background(cell, "F1F5F9")  # Slate-100
    ip = cell.paragraphs[0]
    irun = ip.add_run(
        "💡 INSTRUCTIONS FOR TENANT:\n"
        "1. Fill in the fields and tables below with your company's actual data.\n"
        "2. Do not modify the section header titles (e.g. [SECTION: BRAND & PROFILE]) as our AI parser uses them to bifurcate your data.\n"
        "3. Once filled, save this document and upload it via the Admin Console. The AI engine will parse your menu, catalog, and FAQs into your live bot!"
    )
    irun.font.size = Pt(9.5)
    irun.font.color.rgb = RGBColor(51, 65, 85)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 1: BRAND & PROFILE
    # ──────────────────────────────────────────────────────────────────────────
    h1 = doc.add_heading("[SECTION: BRAND & PROFILE]", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    brand_table = doc.add_table(rows=11, cols=3)
    brand_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(brand_table)

    headers = ["Parameter", "Description & Constraints", "Your Company Input (Fill Here)"]
    for i, h in enumerate(headers):
        c = brand_table.cell(0, i)
        _set_cell_background(c, "1E293B")
        cp = c.paragraphs[0]
        crun = cp.add_run(h)
        crun.font.bold = True
        crun.font.size = Pt(9.5)
        crun.font.color.rgb = RGBColor(255, 255, 255)

    brand_fields = [
        ("Company Legal Name", "Official registered company or entity name", config["default_company"]),
        ("Brand / Display Name", "Public name displayed to WhatsApp users", config["default_company"]),
        ("Tenant ID / Code", "Short unique identifier (lowercase, e.g. dlf, skyline)", config["default_tenant"]),
        ("AI Bot Name", "Friendly name of your AI assistant", config["default_bot"]),
        ("Brand Tagline", "1-2 sentence core value proposition", config["tagline"]),
        ("Official Website URL", "Company website link", f"https://www.{config['default_tenant']}.com"),
        ("Customer Support Email", "Official support or sales contact email", f"support@{config['default_tenant']}.com"),
        ("Support Phone Number", "Official contact telephone with country code", "+91 9876543210"),
        ("WhatsApp Bot Number", "Dedicated phone number to connect to WhatsApp bot", "+91 XXXXXXXXXX"),
        ("Business Working Hours", "Daily operational hours (e.g. 09:30 AM - 06:30 PM)", config["default_hours"]),
    ]

    for row_idx, (f_name, f_desc, f_val) in enumerate(brand_fields, start=1):
        c0 = brand_table.cell(row_idx, 0)
        c1 = brand_table.cell(row_idx, 1)
        c2 = brand_table.cell(row_idx, 2)

        _set_cell_background(c0, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
        _set_cell_background(c1, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
        _set_cell_background(c2, "FEF3C7" if row_idx % 2 == 0 else "FFFBEB")  # Soft yellow highlight for input

        p0 = c0.paragraphs[0].add_run(f_name)
        p0.font.bold = True
        p0.font.size = Pt(9)

        p1 = c1.paragraphs[0].add_run(f_desc)
        p1.font.size = Pt(8.5)
        p1.font.italic = True

        p2 = c2.paragraphs[0].add_run(f_val)
        p2.font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 2: WHATSAPP MENU
    # ──────────────────────────────────────────────────────────────────────────
    h2 = doc.add_heading("[SECTION: WHATSAPP MENU]", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(4)

    menu_note = doc.add_paragraph()
    mn_run = menu_note.add_run(
        "CRITICAL WHATSAPP INTERACTIVE BUTTON CONSTRAINTS:\n"
        "• Title: Max 24 characters (e.g., 'Explore Properties', 'Track Order')\n"
        "• Description: Max 72 characters (e.g., '2, 3 & 4 BHK apartments in Prime Whitefield')\n"
        "• Section: Max 24 characters (e.g., 'Properties', 'Support')"
    )
    mn_run.font.size = Pt(8.5)
    mn_run.font.color.rgb = RGBColor(180, 83, 9)  # Amber-700

    menu_table = doc.add_table(rows=len(config["menu_items"]) + 1, cols=4)
    menu_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(menu_table)

    m_headers = ["Button Title (<=24 chars)", "Description (<=72 chars)", "Category Section", "Action Mapping"]
    for i, h in enumerate(m_headers):
        c = menu_table.cell(0, i)
        _set_cell_background(c, "1E293B")
        cp = c.paragraphs[0].add_run(h)
        cp.font.bold = True
        cp.font.size = Pt(9)
        cp.font.color.rgb = RGBColor(255, 255, 255)

    for row_idx, (m_title, m_desc, m_sec, m_act) in enumerate(config["menu_items"], start=1):
        for col_idx, val in enumerate([m_title, m_desc, m_sec, m_act]):
            cell = menu_table.cell(row_idx, col_idx)
            _set_cell_background(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            prun = cell.paragraphs[0].add_run(val)
            prun.font.size = Pt(8.5)
            if col_idx == 0:
                prun.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 3: CATALOG & OFFERINGS
    # ──────────────────────────────────────────────────────────────────────────
    h3 = doc.add_heading("[SECTION: CATALOG & OFFERINGS]", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(4)

    cat_cols = config["catalog_columns"]
    cat_rows = config["catalog_rows"]

    cat_table = doc.add_table(rows=len(cat_rows) + 3, cols=len(cat_cols))
    cat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(cat_table)

    for i, h in enumerate(cat_cols):
        c = cat_table.cell(0, i)
        _set_cell_background(c, "0F766E")  # Teal-700
        cp = c.paragraphs[0].add_run(h)
        cp.font.bold = True
        cp.font.size = Pt(8)
        cp.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_data in enumerate(cat_rows, start=1):
        for c_idx, val in enumerate(row_data):
            cell = cat_table.cell(r_idx, c_idx)
            _set_cell_background(cell, "F0FDFA" if r_idx % 2 == 0 else "FFFFFF")
            prun = cell.paragraphs[0].add_run(str(val))
            prun.font.size = Pt(8)

    # Add 2 blank rows for user addition
    for b_idx in range(len(cat_rows) + 1, len(cat_rows) + 3):
        for c_idx in range(len(cat_cols)):
            cell = cat_table.cell(b_idx, c_idx)
            _set_cell_background(cell, "FEF3C7")
            prun = cell.paragraphs[0].add_run("[Enter Item]" if c_idx == 0 else "")
            prun.font.size = Pt(8)
            prun.font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 4: FAQS & POLICIES
    # ──────────────────────────────────────────────────────────────────────────
    h4 = doc.add_heading("[SECTION: FAQS & POLICIES]", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(4)

    faq_table = doc.add_table(rows=len(config["faq_rows"]) + 3, cols=3)
    faq_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(faq_table)

    f_headers = ["Frequently Asked Customer Question", "Verified Official Answer", "Topic / Category"]
    for i, h in enumerate(f_headers):
        c = faq_table.cell(0, i)
        _set_cell_background(c, "4338CA")  # Indigo-700
        cp = c.paragraphs[0].add_run(h)
        cp.font.bold = True
        cp.font.size = Pt(9)
        cp.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, (q, a, cat) in enumerate(config["faq_rows"], start=1):
        for c_idx, val in enumerate([q, a, cat]):
            cell = faq_table.cell(r_idx, c_idx)
            _set_cell_background(cell, "EEF2FF" if r_idx % 2 == 0 else "FFFFFF")
            prun = cell.paragraphs[0].add_run(val)
            prun.font.size = Pt(8.5)
            if c_idx == 0:
                prun.font.bold = True

    for b_idx in range(len(config["faq_rows"]) + 1, len(config["faq_rows"]) + 3):
        for c_idx in range(3):
            cell = faq_table.cell(b_idx, c_idx)
            _set_cell_background(cell, "FEF3C7")
            prun = cell.paragraphs[0].add_run("[Add Question]" if c_idx == 0 else "")
            prun.font.size = Pt(8.5)
            prun.font.italic = True

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ──────────────────────────────────────────────────────────────────────────
    # SECTION 5: BOT PERSONA & GUARDRAILS
    # ──────────────────────────────────────────────────────────────────────────
    h5 = doc.add_heading("[SECTION: BOT PERSONA & GUARDRAILS]", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(4)

    guard_table = doc.add_table(rows=5, cols=2)
    guard_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _set_table_borders(guard_table)

    g_headers = ["Configuration Parameter", "Bot Behavior / Value"]
    for i, h in enumerate(g_headers):
        c = guard_table.cell(0, i)
        _set_cell_background(c, "991B1B")  # Red-800
        cp = c.paragraphs[0].add_run(h)
        cp.font.bold = True
        cp.font.size = Pt(9)
        cp.font.color.rgb = RGBColor(255, 255, 255)

    guard_rows = [
        ("Tone & Persona", "Polite, consultative, crisp, and brand-aligned"),
        ("Welcome Message", f"Welcome to {config['default_company']}! How can I assist you today?"),
        ("Strict Guardrails (Never-Do List)", "1. Never offer unauthorized discounts or unlisted pricing.\n2. Never promise dates or legal commitments outside official brochure.\n3. Never request passwords, OTPs, or credit card details."),
        ("Human Escalation Trigger", "Escalate to live sales agent when user asks for 'human', 'agent', 'manager', or expresses grievance."),
    ]

    for r_idx, (param, val) in enumerate(guard_rows, start=1):
        c0 = guard_table.cell(r_idx, 0)
        c1 = guard_table.cell(r_idx, 1)
        _set_cell_background(c0, "FEF2F2" if r_idx % 2 == 0 else "FFFFFF")
        _set_cell_background(c1, "FEF2F2" if r_idx % 2 == 0 else "FFFFFF")
        p0 = c0.paragraphs[0].add_run(param)
        p0.font.bold = True
        p0.font.size = Pt(9)
        p1 = c1.paragraphs[0].add_run(val)
        p1.font.size = Pt(8.5)

    # Save DOCX
    doc.save(str(output_path))
    return output_path


# ──────────────────────────────────────────────────────────────────────────────
# 2. PDF Generation Helper Functions (PyMuPDF / fitz)
# ──────────────────────────────────────────────────────────────────────────────

class PDFKitWriter:
    """Creates a high quality, paginated PDF Onboarding Kit using PyMuPDF."""

    def __init__(self, vertical_key: str, config: Dict[str, Any], output_path: Path):
        self.vertical_key = vertical_key
        self.config = config
        self.output_path = output_path
        self.doc = fitz.open()
        self.page_w = 595.0
        self.page_h = 842.0
        self.margin = 40.0
        self.content_w = self.page_w - 2 * self.margin
        self.page_num = 0
        self.page = None
        self.y = self.margin
        self._new_page()

    def _new_page(self):
        if self.page:
            self._draw_footer()
        self.page_num += 1
        self.page = self.doc.new_page(width=self.page_w, height=self.page_h)
        self.y = self.margin
        self._draw_header()

    def _draw_header(self):
        # Subtle header rule
        self.page.insert_text(
            (self.margin, self.y + 10),
            f"AI WhatsApp Chatbot Onboarding Kit — {self.config['title']}",
            fontsize=8,
            fontname="helv",
            color=(0.4, 0.45, 0.5),
        )
        self.page.draw_line(
            (self.margin, self.y + 16),
            (self.page_w - self.margin, self.y + 16),
            color=(0.85, 0.88, 0.92),
            width=0.7,
        )
        self.y += 26

    def _draw_footer(self):
        self.page.draw_line(
            (self.margin, self.page_h - 30),
            (self.page_w - self.margin, self.page_h - 30),
            color=(0.85, 0.88, 0.92),
            width=0.7,
        )
        self.page.insert_text(
            (self.margin, self.page_h - 18),
            "Confidential — Fill & upload to configure your tenant bot automatically",
            fontsize=7.5,
            fontname="helv",
            color=(0.5, 0.55, 0.6),
        )
        self.page.insert_text(
            (self.page_w - self.margin - 35, self.page_h - 18),
            f"Page {self.page_num}",
            fontsize=8,
            fontname="hebo",
            color=(0.2, 0.25, 0.3),
        )

    def _ensure_space(self, needed_h: float):
        if self.y + needed_h > self.page_h - 45:
            self._new_page()

    def draw_section_header(self, title: str, color=(0.12, 0.16, 0.23)):
        self._ensure_space(32)
        # Background pill
        rect = fitz.Rect(self.margin, self.y, self.page_w - self.margin, self.y + 20)
        self.page.draw_rect(rect, color=color, fill=color)
        self.page.insert_text(
            (self.margin + 8, self.y + 14),
            title,
            fontsize=9.5,
            fontname="hebo",
            color=(1.0, 1.0, 1.0),
        )
        self.y += 26

    def draw_table(
        self,
        headers: List[str],
        rows: List[List[str]],
        col_widths: List[float],
        header_bg=(0.12, 0.16, 0.23),
        font_size=7.2,
    ):
        line_height = font_size * 1.35
        pad = 4.0

        # Draw Header Row
        self._ensure_space(22)
        h_rect = fitz.Rect(self.margin, self.y, self.page_w - self.margin, self.y + 18)
        self.page.draw_rect(h_rect, color=header_bg, fill=header_bg)

        # Draw vertical lines for header cells
        x_grid = self.margin
        for w in col_widths:
            self.page.draw_line((x_grid, self.y), (x_grid, self.y + 18), color=(0.7, 0.7, 0.7), width=0.5)
            x_grid += w
        self.page.draw_line((self.page_w - self.margin, self.y), (self.page_w - self.margin, self.y + 18), color=(0.7, 0.7, 0.7), width=0.5)

        x_curr = self.margin
        for idx, h_text in enumerate(headers):
            w = col_widths[idx]
            self.page.insert_text(
                (x_curr + pad, self.y + 12),
                h_text[:28],
                fontsize=font_size + 0.3,
                fontname="hebo",
                color=(1.0, 1.0, 1.0),
            )
            x_curr += w
        self.y += 18

        # Draw Data Rows
        for r_idx, row in enumerate(rows):
            # Calculate needed height for wrapped cells
            max_lines = 1
            wrapped_cells = []
            for c_idx, cell_text in enumerate(row):
                w = col_widths[c_idx] - (2 * pad)
                # Simple character-based wrap estimation
                approx_chars_per_line = max(10, int(w / (font_size * 0.52)))
                words = str(cell_text).split()
                lines = []
                cur_line = ""
                for wd in words:
                    if len(cur_line) + len(wd) + 1 <= approx_chars_per_line:
                        cur_line = f"{cur_line} {wd}".strip()
                    else:
                        if cur_line:
                            lines.append(cur_line)
                        cur_line = wd
                if cur_line:
                    lines.append(cur_line)
                lines = lines or [""]
                max_lines = max(max_lines, len(lines))
                wrapped_cells.append(lines)

            row_h = max_lines * line_height + (2 * pad)
            self._ensure_space(row_h)

            bg_color = (0.97, 0.98, 0.99) if r_idx % 2 == 0 else (1.0, 1.0, 1.0)
            row_rect = fitz.Rect(self.margin, self.y, self.page_w - self.margin, self.y + row_h)
            self.page.draw_rect(row_rect, color=(0.75, 0.78, 0.82), fill=bg_color, width=0.5)

            # Draw vertical column lines for table cells
            x_grid = self.margin
            for w in col_widths:
                self.page.draw_line((x_grid, self.y), (x_grid, self.y + row_h), color=(0.75, 0.78, 0.82), width=0.5)
                x_grid += w
            self.page.draw_line((self.page_w - self.margin, self.y), (self.page_w - self.margin, self.y + row_h), color=(0.75, 0.78, 0.82), width=0.5)

            x_curr = self.margin
            for c_idx, lines in enumerate(wrapped_cells):
                w = col_widths[c_idx]
                for l_idx, line_str in enumerate(lines):
                    y_pos = self.y + pad + (l_idx + 1) * line_height - 2
                    is_bold = (c_idx == 0 and r_idx < len(rows))
                    self.page.insert_text(
                        (x_curr + pad, y_pos),
                        line_str,
                        fontsize=font_size,
                        fontname="hebo" if is_bold else "helv",
                        color=(0.1, 0.15, 0.2),
                    )
                x_curr += w
            self.y += row_h

        self.y += 8

    def build(self):
        # Title & Intro
        self._ensure_space(55)
        self.page.insert_text(
            (self.margin, self.y + 14),
            f"AI WhatsApp Chatbot Onboarding Guide",
            fontsize=15,
            fontname="hebo",
            color=(0.12, 0.16, 0.23),
        )
        self.page.insert_text(
            (self.margin, self.y + 28),
            f"Vertical: {self.config['title']} ({self.config['slug']})",
            fontsize=10,
            fontname="hebo",
            color=(0.06, 0.46, 0.43),  # Teal
        )
        self.page.insert_text(
            (self.margin, self.y + 40),
            "Complete this template and upload it via the Admin Console to automatically configure your bot.",
            fontsize=8.5,
            fontname="helv",
            color=(0.4, 0.45, 0.5),
        )
        self.y += 50

        # Section 1: Brand & Profile
        self.draw_section_header("[SECTION: BRAND & PROFILE]", color=(0.12, 0.16, 0.23))
        brand_rows = [
            ["Company Legal Name", "Official registered entity name", self.config["default_company"]],
            ["Brand Display Name", "Name displayed on WhatsApp", self.config["default_company"]],
            ["Tenant ID / Code", "Short code (e.g. skyline, dlf)", self.config["default_tenant"]],
            ["AI Bot Assistant Name", "Bot assistant persona name", self.config["default_bot"]],
            ["Brand Tagline", "Value proposition sentence", self.config["tagline"]],
            ["Official Website", "Company website URL", f"https://www.{self.config['default_tenant']}.com"],
            ["Support Email", "Support desk inbox", f"support@{self.config['default_tenant']}.com"],
            ["Support Phone", "Helpline number with country code", "+91 9876543210"],
            ["Working Hours", "Operating schedule & days", self.config["default_hours"]],
        ]
        self.draw_table(
            headers=["Parameter", "Description", "Value / Client Input"],
            rows=brand_rows,
            col_widths=[140.0, 175.0, 200.0],
            header_bg=(0.12, 0.16, 0.23),
            font_size=7.2,
        )

        # Section 2: WhatsApp Menu
        self.draw_section_header("[SECTION: WHATSAPP MENU]", color=(0.12, 0.16, 0.23))
        menu_rows = [
            [title, desc, sec, act]
            for (title, desc, sec, act) in self.config["menu_items"]
        ]
        self.draw_table(
            headers=["Button Title (<=24ch)", "Description (<=72ch)", "Section", "Action"],
            rows=menu_rows,
            col_widths=[125.0, 210.0, 90.0, 90.0],
            header_bg=(0.12, 0.16, 0.23),
            font_size=7.0,
        )

        # Section 3: Catalog & Offerings
        self.draw_section_header("[SECTION: CATALOG & OFFERINGS]", color=(0.06, 0.46, 0.43))
        cat_cols = self.config["catalog_columns"]
        # Distribute available 515 pt width among columns
        w_per_col = float(self.content_w) / len(cat_cols)
        col_widths = [w_per_col] * len(cat_cols)

        str_cat_rows = [[str(item) for item in row] for row in self.config["catalog_rows"]]
        self.draw_table(
            headers=cat_cols,
            rows=str_cat_rows,
            col_widths=col_widths,
            header_bg=(0.06, 0.46, 0.43),
            font_size=6.5,
        )

        # Section 4: FAQs & Policies
        self.draw_section_header("[SECTION: FAQS & POLICIES]", color=(0.26, 0.22, 0.79))
        faq_rows = [
            [q, a, cat] for (q, a, cat) in self.config["faq_rows"]
        ]
        self.draw_table(
            headers=["Customer Question", "Official Verified Answer", "Category"],
            rows=faq_rows,
            col_widths=[160.0, 275.0, 80.0],
            header_bg=(0.26, 0.22, 0.79),
            font_size=7.0,
        )

        # Section 5: Bot Persona & Guardrails
        self.draw_section_header("[SECTION: BOT PERSONA & GUARDRAILS]", color=(0.6, 0.1, 0.1))
        guard_rows = [
            ["Tone & Style", "Consultative, professional and brand-aligned"],
            ["Welcome Greeting", f"Welcome to {self.config['default_company']}! How can I assist you today?"],
            ["Strict Guardrails", "Never offer unauthorized discounts; Never promise unlisted dates; Never ask for OTPs"],
            ["Human Escalation", "Transfer to human sales advisor when user asks for 'human', 'agent' or expresses dispute"],
        ]
        self.draw_table(
            headers=["Parameter", "Rule & Guideline"],
            rows=guard_rows,
            col_widths=[140.0, 375.0],
            header_bg=(0.6, 0.1, 0.1),
            font_size=7.0,
        )

        # Final footer & save
        self._draw_footer()
        self.doc.set_metadata({
            "title": f"Onboarding Kit - {self.config['title']}",
            "subject": "WhatsApp Chatbot Onboarding Template",
            "author": "Antigravity Chatbot Platform",
        })
        self.doc.save(str(self.output_path))
        self.doc.close()
        return self.output_path


def generate_all_kits() -> Dict[str, Dict[str, str]]:
    """Generate both DOCX and PDF kits for all 8 verticals."""
    results = {}
    print(f"Generating Onboarding Kits into: {OUT_DIR}")

    for vertical_key, cfg in VERTICAL_CONFIGS.items():
        docx_path = OUT_DIR / f"{vertical_key}_onboarding_kit.docx"
        pdf_path = OUT_DIR / f"{vertical_key}_onboarding_kit.pdf"

        # Generate DOCX
        generate_docx_kit(vertical_key, cfg, docx_path)

        # Generate PDF
        pdf_writer = PDFKitWriter(vertical_key, cfg, pdf_path)
        pdf_writer.build()

        results[vertical_key] = {
            "title": cfg["title"],
            "slug": vertical_key,
            "docx": str(docx_path),
            "pdf": str(pdf_path),
        }
        print(f"  [OK] Generated {vertical_key.upper()}:")
        print(f"       -> DOCX: {docx_path.name} ({os.path.getsize(docx_path):,} bytes)")
        print(f"       -> PDF:  {pdf_path.name} ({os.path.getsize(pdf_path):,} bytes)")

    print(f"All 8 verticals generated successfully!")
    return results


if __name__ == "__main__":
    generate_all_kits()
